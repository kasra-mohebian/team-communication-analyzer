"""
make_slides.py — build the teammate-facing presentation as a .pptx deck.

A self-contained, image-rich walkthrough of the whole project: the goal, how we
went from unlabelled chats to labels, the data, the two feature views, every
decision (kept / fixed / dropped), the failures, the augmentation, the results,
and the new dashboard — with the scientific reference behind each choice.

    python src/make_slides.py   →   docs/PRESENTATION.pptx

Run make_report_assets.py first so docs/assets/*.png exist.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import config

DOCS = config.ROOT / "docs"
ASSETS = DOCS / "assets"
OUT = str(DOCS / "PRESENTATION.pptx")

# ── palette ──────────────────────────────────────────────────────────────────
NAVY   = RGBColor(0x1B, 0x2A, 0x41)
BLUE   = RGBColor(0x2E, 0x86, 0xC1)
GREEN  = RGBColor(0x27, 0xAE, 0x60)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
RED    = RGBColor(0xC0, 0x39, 0x2B)
GREY   = RGBColor(0x5D, 0x6D, 0x7E)
LIGHT  = RGBColor(0xF4, 0xF6, 0xF7)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x22, 0x22, 0x22)
ARCHC  = {"Bee": RGBColor(0xF4,0xC4,0x30), "Ant": RGBColor(0xC8,0x77,0x2E),
          "Butterfly": RGBColor(0x9B,0x59,0xB6), "Capybara": RGBColor(0x2E,0xCC,0x71),
          "Leech": RGBColor(0xE7,0x4C,0x3C)}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
FONT = "Calibri"


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, fill, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=4, line_spacing=1.05):
    """runs: list of paragraphs; each paragraph is list of (txt, size, bold, color)."""
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2); tf.margin_top = tf.margin_bottom = Pt(2)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line_spacing
        for (txt, size, bold, color) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
            r.font.name = FONT
    return tb


def header(s, kicker, title, tcolor=NAVY):
    rect(s, 0, 0, SW, Inches(0.18), BLUE)
    text(s, Inches(0.6), Inches(0.42), Inches(12.1), Inches(0.4),
         [[(kicker.upper(), 12, True, BLUE)]])
    text(s, Inches(0.6), Inches(0.72), Inches(12.1), Inches(0.9),
         [[(title, 30, True, tcolor)]])
    ln = rect(s, Inches(0.6), Inches(1.62), Inches(12.1), Pt(2), RGBColor(0xD5,0xD8,0xDC))


def pic_fit(s, path, x, y, max_w, max_h):
    iw, ih = Image.open(str(path)).size
    ar = iw / ih
    w, h = max_w, int(max_w / ar)
    if h > max_h:
        h = max_h; w = int(max_h * ar)
    xoff = x + (max_w - w) // 2
    s.shapes.add_picture(str(path), xoff, y, width=w, height=h)
    return w, h


def footer(s, n):
    text(s, Inches(0.6), Inches(7.05), Inches(9), Inches(0.35),
         [[("SocialCompass — Team Communication Analyzer", 9, False, GREY)]])
    text(s, Inches(11.8), Inches(7.05), Inches(1.0), Inches(0.35),
         [[(str(n), 9, False, GREY)]], align=PP_ALIGN.RIGHT)


def bullets(s, x, y, w, items, size=15, gap=8, color=DARK):
    runs = []
    for it in items:
        if isinstance(it, tuple):  # (text, indent_level)
            txt, lvl = it
        else:
            txt, lvl = it, 0
        bullet = "—  " if lvl else "•  "
        runs.append([(bullet + txt, size - (1 if lvl else 0), False, GREY if lvl else color)])
    text(s, x, y, w, Inches(5), runs, space_after=gap, line_spacing=1.08)


N = 0
def number():
    global N; N += 1; return N


# ═════════════════════════════════════════════════════════════════════════════
# 1 · TITLE
# ═════════════════════════════════════════════════════════════════════════════
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(5.0), SW, Inches(0.06), BLUE)
text(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(1.4),
     [[("SocialCompass", 54, True, WHITE)]])
text(s, Inches(0.95), Inches(3.1), Inches(11.5), Inches(0.7),
     [[("Team Communication Analyzer · COIN Archetypes", 22, False, RGBColor(0x9F,0xC5,0xE8))]])
text(s, Inches(0.95), Inches(4.0), Inches(11.0), Inches(1.0),
     [[("From unlabelled group chats to a validated archetype classifier — "
        "the whole journey, the decisions, and what we learned.", 15, False,
        RGBColor(0xBD,0xC3,0xC7))]])
text(s, Inches(0.95), Inches(6.4), Inches(11), Inches(0.5),
     [[("Social Network Analysis seminar  ·  team walkthrough", 13, False,
        RGBColor(0x88,0x95,0xA7))]])

# ═════════════════════════════════════════════════════════════════════════════
# 2 · THE GOAL — 5 archetypes
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "The goal", "Which collaboration role does each person play?")
text(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(0.6),
     [[("Read a team's everyday chat and place each member on the five COIN "
        "collaboration archetypes (Gloor).", 16, False, DARK)]])
arch = [("Bee", "Central connector — starts threads, asks questions, links people"),
        ("Ant", "Reliable task worker — steady, substantive delivery"),
        ("Butterfly", "Information bridge — shares links & news, connects topics"),
        ("Capybara", "Warm supporter — gratitude, encouragement, harmony"),
        ("Leech", "Low-reciprocity — receives more than contributes")]
y = Inches(2.6)
for name, desc in arch:
    rect(s, Inches(0.6), y, Inches(0.28), Inches(0.7), ARCHC[name])
    text(s, Inches(1.05), y, Inches(2.3), Inches(0.7),
         [[(name, 18, True, NAVY)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(3.3), y, Inches(9.4), Inches(0.7),
         [[(desc, 15, False, DARK)]], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.82)
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 3 · THE THREE CHALLENGES
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Why it was hard", "Three problems that shaped every decision")
cards = [("No labels", "Our raw chats had no archetype labels at all — and hand-"
          "labelling thousands of users is infeasible and subjective.", BLUE),
         ("Tiny team", "The team we care about is only 5 people, where network "
          "metrics become unreliable or mathematically degenerate.", ORANGE),
         ("Rare roles", "Butterfly & Capybara barely appear in technical chats, "
          "so a model simply cannot learn them from real data alone.", RED)]
x = Inches(0.6)
for title, body, c in cards:
    rect(s, x, Inches(2.3), Inches(3.9), Inches(0.12), c)
    rect(s, x, Inches(2.42), Inches(3.9), Inches(3.1), LIGHT)
    text(s, x + Inches(0.25), Inches(2.65), Inches(3.4), Inches(0.6),
         [[(title, 20, True, c)]])
    text(s, x + Inches(0.25), Inches(3.35), Inches(3.4), Inches(2.0),
         [[(body, 15, False, DARK)]], line_spacing=1.12)
    x += Inches(4.15)
text(s, Inches(0.6), Inches(5.9), Inches(12.1), Inches(0.6),
     [[("The rest of this story is how we solved all three — honestly.", 16, True, NAVY)]])
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 4 · JOURNEY OVERVIEW
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "The journey", "Every step — and whether it survived")
pic_fit(s, ASSETS / "journey.png", Inches(0.4), Inches(2.2), Inches(12.5), Inches(3.6))
text(s, Inches(0.6), Inches(6.1), Inches(12.1), Inches(0.7),
     [[("Nine major steps: six kept, one repaired, two deliberately dropped. "
        "We walk through each next.", 15, False, GREY)]])
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 5 · FROM NO LABELS TO LABELS
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Step 1", "From no labels to labels — without cheating")
bullets(s, Inches(0.6), Inches(1.9), Inches(7.0), [
    "Weak supervision: a language model reads each user's raw messages and "
    "returns a SOFT label — a spread over the five roles, not one hard class.",
    "Soft labels capture reality: most people blend roles (e.g. 60% Bee / 25% Ant).",
    "Not circular — the key guarantee:",
    ("Inputs X = measured behaviour + language features.", 1),
    ("Labels y = produced independently from the raw text.", 1),
    ("X and y come from different sources → agreement measures validity, "
     "not self-consistency.", 1),
], size=15, gap=9)
rect(s, Inches(8.0), Inches(2.1), Inches(4.7), Inches(3.6), LIGHT)
text(s, Inches(8.25), Inches(2.35), Inches(4.2), Inches(0.5),
     [[("Why soft labels?", 17, True, NAVY)]])
text(s, Inches(8.25), Inches(3.0), Inches(4.2), Inches(2.6),
     [[("A hard \"Bee\" label scores an honest 27% Bee / 26% Ant person as simply "
        "wrong. Soft labels keep that nuance — and let us measure how much "
        "probability mass the model shares with the label.", 14, False, DARK)]],
     line_spacing=1.15)
text(s, Inches(0.6), Inches(6.55), Inches(12.1), Inches(0.5),
     [[("Basis: Ratner et al. (2020) weak supervision · Hinton et al. (2015) soft-label distillation.",
        11, False, GREY)]])
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 6 · DATA
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Step 2", "The data — train on public chats, test on ours")
pic_fit(s, ASSETS / "datasets.png", Inches(0.5), Inches(1.9), Inches(7.6), Inches(4.2))
bullets(s, Inches(8.3), Inches(2.1), Inches(4.5), [
    "300 labelled users for training.",
    "Slack + Nankani + Mbada = real public corpora.",
    "Synthetic = disclosed in-domain augmentation.",
    "Our 5-person team is TEST ONLY — never trained, tuned, or labelled for training.",
], size=15, gap=11)
text(s, Inches(8.3), Inches(5.4), Inches(4.5), Inches(1.0),
     [[("Strict train/test separation is what makes the final number honest.",
        14, True, NAVY)]])
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 7 · TWO VIEWS (architecture)
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Step 3", "Two independent views of each person")
pic_fit(s, DOCS / "architecture.png", Inches(0.4), Inches(1.85), Inches(8.6), Inches(4.9))
bullets(s, Inches(9.2), Inches(2.1), Inches(3.7), [
    "View A — 13 network & behaviour metrics (audited & repaired in v2).",
    "View B — 17 grounded text-style signals (scale-invariant).",
    "Each is weak alone on 5 people; together they are robust.",
    "Late fusion combines both into one soft distribution.",
], size=14, gap=10)
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 8 · TEXT FEATURES TABLE
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Step 3 · detail", "17 text features — every one is citable")
pic_fit(s, DOCS / "text_features_table.png", Inches(0.5), Inches(1.85), Inches(12.3), Inches(4.5))
text(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.5),
     [[("Empath (r=0.906 with LIWC) · VADER sentiment · Searle speech acts · "
        "structural pragmatics. No ad-hoc keyword lists.", 12, False, GREY)]])
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 9 · DECISIONS LOG
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "The decisions", "Kept · Fixed · Dropped")
rows = [("Soft over hard labels", "KEPT", GREEN),
        ("Rule-only classifier (Slide-6 thresholds)", "DROPPED", RED),
        ("v1 interaction metrics (fake reply chain)", "FIXED → v2", ORANGE),
        ("Two-view late fusion", "KEPT", GREEN),
        ("Garimella 6.3M msgs (off-domain spam)", "DROPPED", RED),
        ("Disclosed synthetic augmentation", "KEPT", GREEN),
        ("Grounded, citable text features", "KEPT", GREEN),
        ("Soft accuracy as the headline metric", "KEPT", GREEN),
        ("Balanced 50/50 fusion weight (no leakage)", "KEPT", GREEN)]
y = Inches(1.95)
for label, tag, c in rows:
    rect(s, Inches(0.6), y, Inches(8.7), Inches(0.46), LIGHT)
    text(s, Inches(0.8), y, Inches(8.3), Inches(0.46),
         [[(label, 14, False, DARK)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, Inches(9.5), y, Inches(3.2), Inches(0.46), c)
    text(s, Inches(9.5), y, Inches(3.2), Inches(0.46),
         [[(tag, 13, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.52)
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 10 · WHAT FAILED
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Honesty", "What failed — and what it taught us", tcolor=RED)
fails = [("Garimella (6.3M real messages)",
          "Broadcast forwards, bots, political spam. Lesson: domain match beats raw realness."),
         ("Rule-only classifier",
          "Readable, but thresholds don't transfer across group sizes & can't express blends."),
         ("macOS segfault (two transformer stacks)",
          "Fixed with a style-only path that skips the heavy embedding — also faster."),
         ("First synthetic set too sparse (~4 msgs/user)",
          "Hurt the score; rewrote longer chats. Synthetic must match real density, not just labels.")]
y = Inches(2.0)
for t, b in fails:
    rect(s, Inches(0.6), y, Inches(0.12), Inches(1.0), RED)
    text(s, Inches(0.9), y, Inches(11.8), Inches(0.45), [[(t, 16, True, NAVY)]])
    text(s, Inches(0.9), y + Inches(0.42), Inches(11.8), Inches(0.55), [[(b, 14, False, DARK)]])
    y += Inches(1.15)
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 11 · AUGMENTATION
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Step 4", "Fixing class imbalance — honestly")
pic_fit(s, ASSETS / "ablation.png", Inches(0.5), Inches(1.9), Inches(7.6), Inches(4.3))
bullets(s, Inches(8.3), Inches(2.1), Inches(4.5), [
    "Butterfly & Capybara were unlearnable (F1 ≈ 0).",
    "We added disclosed, in-domain synthetic team chats — generated, then "
    "labelled by the SAME independent process as real users.",
    "Fully disclosed, toggleable, and the blind team test stays 100% real.",
], size=14, gap=10)
text(s, Inches(8.3), Inches(5.2), Inches(4.5), Inches(1.2),
     [[("Result: rare roles become learnable; macro-F1 0.27 → 0.56.", 16, True, GREEN)]])
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 12 · SOFT ACCURACY
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "How we judge it", "Soft accuracy — the honest yardstick")
text(s, Inches(0.6), Inches(2.0), Inches(12.1), Inches(1.2),
     [[("On a fuzzy target, scoring only the top guess as right-or-wrong throws "
        "away most of the truth. Someone who is 27% Bee / 26% Ant is not simply "
        "\"wrong\" if we pick Ant.", 17, False, DARK)]], line_spacing=1.15)
rect(s, Inches(2.5), Inches(3.6), Inches(8.3), Inches(1.2), LIGHT)
text(s, Inches(2.5), Inches(3.6), Inches(8.3), Inches(1.2),
     [[("Soft accuracy  =  mean  Σ min(model, label)", 24, True, NAVY)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(0.6), Inches(5.1), Inches(12.1), Inches(1.2),
     [[("= the probability mass the model shares with the independent label. "
        "It credits partial agreement and is the right metric for blended roles.",
        16, False, DARK)]], line_spacing=1.15)
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 13 · HELD-OUT RESULTS
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Results", "Blind validation on the real team")
pic_fit(s, ASSETS / "holdout.png", Inches(0.5), Inches(1.9), Inches(12.3), Inches(4.3))
text(s, Inches(0.6), Inches(6.45), Inches(12.1), Inches(0.6),
     [[("3 of 5 exact; the two misses are genuine near-ties whose distributions "
        "still overlap the label strongly — exactly what soft accuracy captures.",
        13, False, GREY)]])
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 14 · TEAM RESULT
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Our team", "The result, member by member")
pic_fit(s, ASSETS / "team_snapshot.png", Inches(0.4), Inches(1.9), Inches(12.5), Inches(4.4))
text(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.5),
     [[("Nobody is a single pure type — the model returns each person's realistic "
        "blend of roles.", 13, False, GREY)]])
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 15 · DASHBOARD
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "The product", "A friendly dashboard anyone can read")
pic_fit(s, ASSETS / "team_snapshot.png", Inches(0.5), Inches(1.95), Inches(7.4), Inches(4.3))
bullets(s, Inches(8.1), Inches(2.0), Inches(4.7), [
    "Overview — team make-up, who's-who, role-strength bars, team-health check.",
    "Who talks to whom — interaction network sized by activity, coloured by role.",
    "Activity — message volume over time + everyone-online heatmap.",
    "Member profile — one place per person: role, radar, breakdown, highlights, recent messages.",
], size=14, gap=10)
text(s, Inches(8.1), Inches(5.7), Inches(4.7), Inches(0.8),
     [[("Rewritten to drop jargon and internal machinery — built for non-technical "
        "teammates.", 13, True, NAVY)]])
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 16 · LIMITATIONS
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Honest limits", "What we don't claim")
bullets(s, Inches(0.6), Inches(2.0), Inches(12.1), [
    "Butterfly & Capybara stay the hardest classes, even after augmentation.",
    "Betweenness centrality is near-zero on a 5-person near-clique — a known small-group limit.",
    "Team predictions carry modest confidence (~25–28%) due to the train→team domain shift.",
    "The held-out test is only 5 people — so we never tune to it and headline a distributional metric.",
], size=17, gap=14)
text(s, Inches(0.6), Inches(6.2), Inches(12.1), Inches(0.6),
     [[("Stated, not hidden — the soft metric and disclosed augmentation exist "
        "because we took these seriously.", 15, True, NAVY)]])
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 17 · REFERENCES
# ═════════════════════════════════════════════════════════════════════════════
s = slide(); header(s, "Foundations", "Scientific references")
refs1 = ["Gloor (2017/2022) — COIN archetypes",
         "Ratner et al. (2020) — weak supervision (Snorkel)",
         "Hinton et al. (2015) — soft-label distillation",
         "Breiman (2001) — Random Forests",
         "Fast et al. (2016) — Empath (r=0.906 w/ LIWC)",
         "Hutto & Gilbert (2014) — VADER sentiment",
         "Mohammad & Turney (2013) — NRC EmoLex"]
refs2 = ["Tausczik & Pennebaker (2010) — LIWC",
         "Danescu-Niculescu-Mizil et al. (2013) — Politeness",
         "Searle (1976) — speech acts",
         "Covington & McFall (2010) — MATTR",
         "Freeman (1977) / Brandes (2001) — betweenness",
         "Goh & Barabási (2008) — burstiness",
         "Kullback & Leibler (1951) — KL divergence"]
bullets(s, Inches(0.6), Inches(2.0), Inches(6.1), refs1, size=14, gap=10)
bullets(s, Inches(6.9), Inches(2.0), Inches(6.0), refs2, size=14, gap=10)
footer(s, number())

# ═════════════════════════════════════════════════════════════════════════════
# 18 · CLOSING
# ═════════════════════════════════════════════════════════════════════════════
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(4.4), SW, Inches(0.06), GREEN)
text(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.0),
     [[("Thank you", 46, True, WHITE)]])
text(s, Inches(0.95), Inches(3.5), Inches(11.5), Inches(0.8),
     [[("From unlabelled chats to a validated, honest archetype classifier — "
        "every decision earned its place.", 17, False, RGBColor(0xBD,0xC3,0xC7))]])
text(s, Inches(0.95), Inches(5.0), Inches(11.5), Inches(0.6),
     [[("Soft accuracy 74% · cosine 0.87 · 3/5 exact on a blind 5-person team · "
        "300 labelled training users", 15, True, RGBColor(0x9F,0xC5,0xE8))]])
text(s, Inches(0.95), Inches(6.2), Inches(11.5), Inches(0.5),
     [[("Live demo:  streamlit run src/dashboard.py", 14, False, RGBColor(0x88,0x95,0xA7))]])


prs.save(OUT)
print(f"✅ wrote {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
